"""openpyxl renderer of the timetable layout: one "Orario" sheet."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .layout import EMPTY_BG, HEADER_BG, TITLE_BG, WHITE

SHEET_NAME = "Orario"
COLUMN_WIDTH = 15
PAPER_A4 = 9
FIRST_CLASS_COLUMN = 3  # DATA and Orario come first

_THIN = Side(style="thin", color="FF000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center")


def _fill(color):
    return PatternFill("solid", start_color="FF" + color.lstrip("#").upper())


def _font(color="#000000", **kw):
    return Font(name="Calibri", color="FF" + color.lstrip("#").upper(), **kw)


def render_xlsx(layout):
    """The xlsx bytes of `layout`, styled like the PDF."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    width = len(layout.header)

    for number, line in enumerate(layout.title_lines, start=1):
        ws.merge_cells(
            start_row=number, start_column=1, end_row=number, end_column=width
        )
        cell = ws.cell(row=number, column=1, value=line)
        cell.font = _font(bold=True, size=14)
        cell.alignment = _CENTER
        for col in range(1, width + 1):  # the band colours the merged range
            ws.cell(row=number, column=col).fill = _fill(TITLE_BG)

    header_row = len(layout.title_lines) + 1
    for col, name in enumerate(layout.header, start=1):
        cell = ws.cell(row=header_row, column=col, value=name)
        cell.fill = _fill(HEADER_BG)
        cell.font = _font(WHITE, bold=True)
        cell.alignment = _CENTER
        cell.border = _BORDER

    row = header_row + 1
    for number, block in enumerate(layout.days):
        if number:
            for col in range(1, width + 1):
                ws.cell(row=row, column=col).fill = _fill(EMPTY_BG)
            ws.row_dimensions[row].height = 6
            row += 1
        first = row
        for hour_row in block.rows:
            hour = ws.cell(row=row, column=2, value=hour_row.label)
            hour.alignment = _CENTER
            hour.border = _BORDER
            hour.font = _font()
            for offset, item in enumerate(hour_row.cells):
                cell = ws.cell(
                    row=row, column=FIRST_CLASS_COLUMN + offset, value=item.text or None
                )
                cell.fill = _fill(item.background)
                cell.font = _font(bold=item.bold)
                cell.alignment = _CENTER
                cell.border = _BORDER
            row += 1
        last = row - 1
        if last > first:
            ws.merge_cells(start_row=first, start_column=1, end_row=last, end_column=1)
        for r in range(first, last + 1):
            ws.cell(row=r, column=1).border = _BORDER
            ws.cell(row=r, column=1).fill = _fill(block.background)
        day = ws.cell(row=first, column=1, value=block.label)
        day.font = _font(WHITE, italic=True)
        day.alignment = Alignment(
            horizontal="center", vertical="center", text_rotation=90
        )

    for col in range(1, width + 1):
        ws.column_dimensions[get_column_letter(col)].width = COLUMN_WIDTH
    ws.freeze_panes = ws.cell(row=header_row + 1, column=FIRST_CLASS_COLUMN)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = PAPER_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
