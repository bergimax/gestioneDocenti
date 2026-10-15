import datetime
import io

import pytest
from django.test import Client
from openpyxl import load_workbook

from accounts.models import User
from export.excel import render_xlsx
from export.layout import build_layout
from scheduling import views
from scheduling.models import Slot, Timetable, TimetableState
from school.models import CalendarEntry, Class, School, Subject, Teacher

D = datetime.date
MON = D(2026, 10, 12)
URL = f"/orario/{MON:%Y-%m-%d}/export-excel/"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def login(role):
    user = User.objects.create_user(
        f"{role.lower()}@example.com", "pw-long-enough", role
    )
    client = Client()
    client.force_login(user)
    return client


@pytest.fixture(autouse=True)
def fixed_today(monkeypatch):
    monkeypatch.setattr(views, "today", lambda: D(2026, 10, 14))


@pytest.fixture
def school(db):
    return School.objects.create(
        name="Agenzia Formativa don Angelo Tedoldi",
        formative_year="2026-2027",
        year_start=D(2026, 9, 1),
        year_end=D(2027, 6, 30),
    )


@pytest.fixture
def world(school):
    class W:
        pass

    w = W()
    w.school = school
    w.itc = Class.objects.create(school=school, name="I I.T.C.")
    w.zeta = Class.objects.create(school=school, name="Zeta")
    w.el2 = Class.objects.create(school=school, name="II ELETTRICISTI")
    w.op1 = Class.objects.create(school=school, name="I OP. INFORM.")
    w.rossi = Teacher.objects.create(
        school=school, name="ROSSI MARIO", type="ASSUNTO", first_time=False
    )
    w.mat = Subject.objects.create(
        school=school, name="Matematica", kind="TEORIA", cognitive_weight="MEDIO"
    )
    w.tt = Timetable.objects.create(school=school, week_start=MON, version=5)
    return w


def add(w, klass, day="LUNEDI", hour=8, **kw):
    return Slot.objects.create(
        timetable=w.tt,
        school_class=klass,
        teacher=w.rossi,
        subject=w.mat,
        day=day,
        start_hour=hour,
        end_hour=hour + 1,
        **kw,
    )


def sheet(data):
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Orario"]
    return wb["Orario"]


def rgb(cell):
    return "#" + cell.fill.start_color.rgb[2:].lower()


def populate(w):
    add(w, w.op1, hour=8)
    add(w, w.itc, hour=9)
    add(w, w.el2, day="MARTEDI", hour=10)
    add(w, w.op1, day="VENERDI", hour=8, paired=True, paired_class=w.el2)
    add(w, w.el2, day="VENERDI", hour=8, paired=True, paired_class=w.op1)
    CalendarEntry.objects.create(
        school=w.school, date=MON, stage_class=w.zeta, group=None
    )


def test_cells_and_fills_equal_the_layout_model(world):
    populate(world)
    layout = build_layout(world.school, MON, world.tt)
    ws = sheet(render_xlsx(layout))
    assert ws["A1"].value == layout.title_lines[0]
    assert ws["A2"].value == "Orario scolastico dal 12 al 16 Ottobre"
    for r in (1, 2):
        assert rgb(ws.cell(r, 1)) == "#95b3d7"
        assert ws.cell(r, 1).font.bold
        assert ws.cell(r, 1).font.color.rgb == "FF000000"
    assert {str(m) for m in ws.merged_cells.ranges} >= {"A1:F1", "A2:F2"}

    assert [c.value for c in ws[3]] == list(layout.header)
    assert [c.value for c in ws[3]][:3] == ["DATA", "Orario", "I OP. INFORM."]
    for c in ws[3]:
        assert rgb(c) == "#c00000"
        assert c.font.bold and c.font.color.rgb == "FFFFFFFF"

    row = 4
    merged = {str(m) for m in ws.merged_cells.ranges}
    for number, block in enumerate(layout.days):
        if number:
            assert all(rgb(c) == "#d9d9d9" for c in ws[row])
            row += 1
        first = row
        for hour_row in block.rows:
            assert ws.cell(row, 2).value == hour_row.label
            for offset, item in enumerate(hour_row.cells):
                cell = ws.cell(row, 3 + offset)
                assert (cell.value or "") == item.text
                assert rgb(cell) == item.background
                assert bool(cell.font.bold) == item.bold
                assert cell.alignment.horizontal == "center"
                assert cell.border.left.style == "thin"
            row += 1
        day = ws.cell(first, 1)
        assert day.value == block.label
        assert rgb(day) == block.background
        assert day.alignment.text_rotation == 90
        assert day.font.italic and day.font.color.rgb == "FFFFFFFF"
        if row - 1 > first:
            assert f"A{first}:A{row - 1}" in merged
    assert row - 1 == ws.max_row


def test_specific_cells(world):
    populate(world)
    ws = sheet(render_xlsx(build_layout(world.school, MON, world.tt)))
    assert ws["C4"].value == "ROSSI" and rgb(ws["C4"]) == "#ddd9c4"
    assert ws["D4"].value is None and rgb(ws["D4"]) == "#d9d9d9"
    assert ws["F4"].value == "STAGE" and rgb(ws["F4"]) == "#ffff00"
    assert rgb(ws["A4"]) == "#e46c0a"
    # Tuesday is the second day and uses the blue
    assert rgb(ws["A7"]) == "#0070c0"
    # Friday paired lesson repeats the surname in both classes, unmerged
    friday = next(
        r for r in range(4, ws.max_row + 1) if ws.cell(r, 1).value == "venerdì"
    )
    assert ws.cell(friday, 3).value == ws.cell(friday, 4).value == "ROSSI"


def test_short_days_have_fewer_rows(world):
    populate(world)
    layout = build_layout(world.school, MON, world.tt)
    counts = [len(d.rows) for d in layout.days]
    assert counts[0] == 2 and counts[4] == 1
    ws = sheet(render_xlsx(layout))
    assert ws.max_row == 3 + sum(counts) + 4  # four separator rows


def test_layout_of_the_sheet(world):
    populate(world)
    ws = sheet(render_xlsx(build_layout(world.school, MON, world.tt)))
    for col in "ABCDEF":
        assert ws.column_dimensions[col].width == 15
    assert ws.freeze_panes == "C4"
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.paperSize == 9
    assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 1
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True


def test_two_month_title(world):
    tt = Timetable.objects.create(school=world.school, week_start=D(2026, 9, 28))
    ws = sheet(render_xlsx(build_layout(world.school, D(2026, 9, 28), tt)))
    assert ws["A2"].value == "Orario scolastico dal 28 Settembre al 2 Ottobre"


def test_zero_slots_still_a_valid_file(world):
    ws = sheet(render_xlsx(build_layout(world.school, MON, world.tt)))
    assert ws["A1"].value == world.school.name
    assert ws["A3"].value == "DATA"
    assert ws.max_row >= 3


def test_no_classes_still_a_valid_file(school):
    tt = Timetable.objects.create(school=school, week_start=MON)
    ws = sheet(render_xlsx(build_layout(school, MON, tt)))
    assert [c.value for c in ws[3]] == ["DATA", "Orario"]
    assert ws["A1"].value == school.name


# ---- view -------------------------------------------------------------------


@pytest.mark.parametrize("role", ["ADMIN", "SEGRETERIA"])
@pytest.mark.parametrize("state", [TimetableState.BOZZA, TimetableState.APPROVATO])
def test_download(world, role, state):
    Timetable.objects.filter(pk=world.tt.pk).update(state=state)
    add(world, world.op1)
    r = login(role).get(URL)
    assert r.status_code == 200
    assert r["Content-Type"] == XLSX
    assert 'filename="orario_2026-10-12.xlsx"' in r["Content-Disposition"]
    assert sheet(r.content)["A3"].value == "DATA"


def test_download_changes_nothing(world):
    from scheduling.models import AuditLog

    assert login("ADMIN").get(URL).status_code == 200
    assert Timetable.objects.get(pk=world.tt.pk).version == 5
    assert AuditLog.objects.count() == 0


def test_anonymous_goes_to_login(world, client):
    r = client.get(URL)
    assert r.status_code == 302 and r.url.startswith("/login/")


def test_other_methods_are_405(world):
    assert login("ADMIN").post(URL).status_code == 405


def test_week_without_timetable_is_404_with_message(school):
    r = login("ADMIN").get(URL)
    assert r.status_code == 404
    assert "Nessun orario per questa settimana" in r.content.decode()


def test_page_shows_both_links_in_order(world):
    html = login("SEGRETERIA").get(f"/orario/?settimana={MON:%Y-%m-%d}").content
    html = html.decode()
    assert html.index("Esporta PDF") < html.index("Esporta Excel")
    assert URL in html
