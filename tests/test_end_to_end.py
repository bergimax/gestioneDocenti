"""One complete journey through the real HTTP surface, no shortcuts to the database:
login, setup import, teacher availability by magic link, generation (the worker job
run in-process), review, approval and export. Only the clock is fixed."""

import datetime
import io
import json

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from openpyxl import load_workbook

from accounts.models import User
from availability import grid as availability_grid
from ledger.models import LedgerEntry
from scheduling import tasks
from scheduling.models import Slot, SolveJob, Timetable, TimetableState
from school.models import Teacher

PW = "correct-horse"
HX = {"HTTP_HX_REQUEST": "true"}
MON = datetime.date(2026, 9, 14)
WEEK = f"{MON:%Y-%m-%d}"

DAYS = [f"2026-09-{d}" for d in (14, 15, 16, 17, 18)]
SETUP = {
    "calendar": "data,ore_max_giornata,flag_chiusura\n"
    + "".join(f"{day},6,false\n" for day in DAYS),
    "teachers": "nome,email,tipo\nROSSI MARIO,,ASSUNTO\nBIANCHI LUCA,,CONTRATTO\n",
    "classes": "nome,n_studenti,gruppo\n1A,20,PRIME\n2A,18,SECONDE\n",
    "subjects": "nome,tipo,peso_cognitivo\nMatematica,TEORIA,ALTO\n"
    "Laboratorio,PRATICA,BASSO\n",
    "assignments": "docente_nome,classe_nome,materia_nome,ore_anno\n"
    "ROSSI MARIO,1A,Matematica,4\nROSSI MARIO,2A,Matematica,4\n"
    "BIANCHI LUCA,1A,Laboratorio,3\nBIANCHI LUCA,2A,Laboratorio,3\n",
    # header only: a CSV pairing has no teacher and the solver refuses it (issue #172)
    "pairings": "classe_a,classe_b,materia_nome,note\n",
}


@pytest.fixture(autouse=True)
def fixed_clock(monkeypatch):
    from scheduling import views

    monkeypatch.setattr(views, "today", lambda: datetime.date(2026, 9, 14))


@pytest.fixture
def browser(db):
    """A logged-out client and an admin account created by the operator."""
    User.objects.create_user("admin@example.com", PW, "ADMIN")
    return Client()


def test_setup_to_approved_timetable_and_export(browser):
    # 1. Login through the real form.
    assert browser.get("/").status_code == 302  # anonymous: sent to login
    login = browser.post("/login/", {"username": "admin@example.com", "password": PW})
    assert login.status_code == 302
    assert browser.get("/").status_code == 200

    # 2. Setup wizard: upload each file, then save.
    for file_type, content in SETUP.items():
        r = browser.post(
            f"/setup/upload/{file_type}/",
            {"file": SimpleUploadedFile(f"{file_type}.csv", content.encode())},
        )
        assert r.status_code == 200, file_type
    saved = browser.post(
        "/setup/save/", {"school_name": "ITS Test", "school_year": "2026-2027"}
    )
    assert saved.status_code in (200, 302), saved.context["blockers"]
    contractor = Teacher.objects.get(name="BIANCHI LUCA")
    assert contractor.type == "CONTRATTO"

    # 3. The contractor is invited with a magic link and fills the grid, anonymous.
    created = browser.post(f"/disponibilita/{WEEK}/docente/{contractor.pk}/link/")
    assert created.status_code == 200
    url = created.context["url"]
    token = url.rstrip("/").rsplit("/", 1)[-1]
    teacher_browser = Client()
    assert teacher_browser.get(f"/l/{token}/").status_code == 200
    grid = availability_grid.default_grid()  # available all week
    saved = teacher_browser.post(f"/l/{token}/salva/", {"grid": json.dumps(grid)}, **HX)
    assert saved.status_code == 200
    assert "_auth_user_id" not in teacher_browser.session

    # 4. Generate: the request queues a job, the worker runs it, the page reports it.
    queued = browser.post(f"/orario/{WEEK}/genera/")
    assert queued.status_code == 200
    job = SolveJob.objects.get()
    tasks.run_job(job.pk)
    job.refresh_from_db()
    assert job.status == "DONE", job.error
    assert browser.get(f"/orario/{WEEK}/stato/").status_code == 200

    # 5. Review: the grid shows lessons and nobody is double-booked.
    timetable = Timetable.objects.get()
    assert timetable.state == TimetableState.BOZZA
    slots = Slot.objects.filter(timetable=timetable)
    assert slots.exists()
    booked = [(s.teacher_id, s.day, s.start_hour) for s in slots]
    assert len(booked) == len(set(booked))
    grid_page = browser.get(f"/orario/{WEEK}/griglia/")
    assert grid_page.status_code == 200
    assert "ROSSI" in grid_page.content.decode()

    # 6. Approve: state moves forward, the ledger records what was delivered.
    approved = browser.post(f"/orario/{WEEK}/approva/", {"confirm": "1"})
    assert approved.status_code == 200
    timetable.refresh_from_db()
    assert timetable.state == TimetableState.APPROVATO
    assert LedgerEntry.objects.filter(timetable=timetable).exists()

    # 7. An approved week refuses further changes and regeneration.
    assert browser.post(f"/orario/{WEEK}/genera/").status_code == 409
    late = teacher_browser.post(f"/l/{token}/salva/", {"grid": json.dumps(grid)}, **HX)
    assert late.status_code == 409

    # 8. Export.
    pdf = browser.get(f"/orario/{WEEK}/export-pdf/")
    assert pdf.status_code == 200
    assert pdf["Content-Type"] == "application/pdf"
    assert (
        b"".join(pdf.streaming_content if pdf.streaming else [pdf.content])[:4]
        == b"%PDF"
    )
    excel = browser.get(f"/orario/{WEEK}/export-excel/")
    assert excel.status_code == 200
    workbook = load_workbook(io.BytesIO(excel.content))
    assert workbook.sheetnames

    # 9. The secretariat can read but not change.
    User.objects.create_user("seg@example.com", PW, "SEGRETERIA")
    seg = Client()
    seg.post("/login/", {"username": "seg@example.com", "password": PW})
    assert seg.get(f"/orario/{WEEK}/griglia/").status_code == 200
    assert seg.post(f"/orario/{WEEK}/genera/").status_code == 403
    assert seg.post(f"/orario/{WEEK}/annulla-approvazione/").status_code == 403
